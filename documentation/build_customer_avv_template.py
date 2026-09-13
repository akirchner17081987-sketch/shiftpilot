from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "SchichtFunk_Kunden_AVV_Vorlage_2026-09-13.docx"

NAVY = "17365D"
LIGHT_BLUE = "EAF1F8"
PALE_BLUE = "F4F8FC"
LIGHT_GRAY = "D9D9D9"
PLACEHOLDER = "FFF2CC"
TEXT_GRAY = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def shade_run(run, fill=PLACEHOLDER):
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    r_pr.append(shd)


def add_text_runs(paragraph, text, bold=False, italic=False):
    parts = re.split(r"(\[\[[^\]]+\]\])", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold
        run.italic = italic
        if part.startswith("[[") and part.endswith("]]" ):
            shade_run(run)
            run.bold = True
    return paragraph


def add_paragraph(doc, text="", *, bold_lead=None, style=None, indent=0, space_after=5):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.12
    if indent:
        paragraph.paragraph_format.left_indent = Inches(indent)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        add_text_runs(paragraph, text[len(bold_lead):])
    else:
        add_text_runs(paragraph, text)
    return paragraph


def add_numbered(doc, number, title, paragraphs):
    heading = doc.add_heading(f"{number} {title}", level=1)
    set_keep_with_next(heading)
    for text in paragraphs:
        add_paragraph(doc, text)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_text_runs(p, text)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[idx])
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text_runs(p, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_document_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(21)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    for name in ("List Bullet", "List Bullet 2"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run("SchichtFunk Kunden AVV Vorlage   |   Version 1.0")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = TEXT_GRAY


doc = Document()
set_document_styles(doc)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
add_footer(section)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(7)
title.add_run("Vereinbarung zur Auftragsverarbeitung")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(3)
r = subtitle.add_run("Vorlage für SchichtFunk Kunden nach Artikel 28 DSGVO")
r.bold = True
r.font.size = Pt(12)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(14)
r = meta.add_run("Version 1.0   |   Stand 13.09.2026")
r.font.size = Pt(9)
r.font.color.rgb = TEXT_GRAY

usage = add_paragraph(
    doc,
    "Verwendungshinweis. Diese Vertragsvorlage ist vor der Unterzeichnung vollständig auszufüllen. Alle gelb markierten Felder, die aktuelle TOM-Fassung, die Unterauftragnehmerliste und die kundenbezogenen Löschfristen müssen geprüft werden. Das Muster ersetzt keine rechtliche Prüfung des konkreten Kundenvertrags.",
    bold_lead="Verwendungshinweis.",
    space_after=10,
)

doc.add_heading("Vertragsparteien", level=1)
add_table(
    doc,
    ["Rolle", "Vertragspartei und Kontakt"],
    [
        (
            "Verantwortlicher\nKunde",
            "[[Vollständiger Firmenname und Rechtsform]]\n[[Anschrift]]\nVertreten durch [[Name und Funktion]]\nDatenschutzkontakt [[E-Mail und Telefon]]",
        ),
        (
            "Auftragsverarbeiter\nSchichtFunk",
            "Alexander Kirchner / SchichtFunk\nNeue Straße 23, 06632 Gröst, Deutschland\nVertreten durch Alexander Kirchner\nDatenschutzkontakt info@schichtfunk.de\nRechtsform bzw. Registerangaben vor Abschluss prüfen und ergänzen: [[Angabe]]",
        ),
    ],
    widths=[1.75, 5.0],
)

add_paragraph(
    doc,
    "Diese Vereinbarung ergänzt den Vertrag über die Nutzung von SchichtFunk vom [[Datum des Hauptvertrags]] (Hauptvertrag). Sie regelt die Verarbeitung personenbezogener Daten durch SchichtFunk im Auftrag des Kunden. Bei Widersprüchen zu Datenschutzfragen geht diese Vereinbarung dem Hauptvertrag vor.",
)

add_numbered(doc, "1", "Gegenstand und Geltungsbereich", [
    "SchichtFunk stellt dem Kunden eine webbasierte Anwendung für Benutzerverwaltung, Mitarbeiterverwaltung, Dienstplanung, Zeiterfassung, Abwesenheiten, Stundenkonten, Personalakten, Benachrichtigungen und vorbereitende Exporte zur Verfügung. Soweit SchichtFunk dabei personenbezogene Daten nach dokumentierten Weisungen des Kunden verarbeitet, handelt SchichtFunk als Auftragsverarbeiter im Sinne von Artikel 28 DSGVO.",
    "Diese Vereinbarung gilt für alle im Hauptvertrag vereinbarten SchichtFunk-Leistungen, bei denen SchichtFunk oder ein genehmigter Unterauftragsverarbeiter auf personenbezogene Daten des Kunden zugreifen kann. Verarbeitungsvorgänge, bei denen SchichtFunk eigene Zwecke und Mittel bestimmt, fallen nur insoweit unter diese Vereinbarung, wie dies gesetzlich vorgesehen ist.",
    "Die Einzelheiten der Verarbeitung ergeben sich aus Anlage 1. Die Anlagen sind Bestandteil dieser Vereinbarung.",
])

add_numbered(doc, "2", "Dauer der Verarbeitung", [
    "Die Verarbeitung beginnt mit der Bereitstellung des Kundenmandanten oder dem ersten Eingang personenbezogener Daten, je nachdem, welches Ereignis früher eintritt. Sie dauert für die Laufzeit des Hauptvertrags und anschließend bis zur vertragsgemäßen Rückgabe oder Löschung der Auftragsdaten.",
    "Gesetzliche Aufbewahrungspflichten und technisch bedingte Backupfenster bleiben unberührt. Daten, die deshalb vorübergehend nicht gelöscht werden dürfen oder können, werden für andere Zwecke gesperrt und nach Ablauf der maßgeblichen Frist gelöscht.",
])

add_numbered(doc, "3", "Weisungen und Verantwortlichkeit des Kunden", [
    "Der Kunde bleibt für die Zulässigkeit der Verarbeitung, die Rechtsgrundlagen, die Erfüllung der Informationspflichten, die Wahrung der Beschäftigtenrechte und die Richtigkeit seiner Weisungen verantwortlich. Dies gilt insbesondere für Krankheits- und Personalaktendaten, Arbeitszeitkontrollen, Auswertungen sowie die Beteiligung eines Betriebsrats oder einer Personalvertretung.",
    "Der Hauptvertrag, diese Vereinbarung, die Konfiguration des Kundenmandanten und dokumentierte Supportaufträge bilden die anfänglichen Weisungen. Weitere Weisungen erteilt der Kunde in Textform über [[vereinbarter Weisungskanal, z. B. info@schichtfunk.de]] durch die in Anlage 1 benannten Personen.",
    "SchichtFunk informiert den Kunden unverzüglich, wenn eine Weisung nach eigener Einschätzung gegen Datenschutzrecht verstößt. SchichtFunk darf die Ausführung bis zur Bestätigung oder Änderung der Weisung aussetzen, soweit dies zum Schutz personenbezogener Daten erforderlich ist.",
    "Der Kunde darf keine Diagnosen, medizinischen Dokumente oder andere besondere Kategorien personenbezogener Daten einstellen, die nicht ausdrücklich in Anlage 1 vereinbart und für den festgelegten Zweck erforderlich sind. Der Status einer Krankheitsabwesenheit ist nur im erforderlichen Umfang zu erfassen.",
])

add_numbered(doc, "4", "Pflichten von SchichtFunk", [
    "SchichtFunk verarbeitet Auftragsdaten ausschließlich nach dokumentierter Weisung des Kunden, soweit keine gesetzliche Verpflichtung zu einer abweichenden Verarbeitung besteht. In diesem Fall informiert SchichtFunk den Kunden vor der Verarbeitung, sofern das Gesetz eine solche Mitteilung nicht untersagt.",
    "SchichtFunk stellt sicher, dass Personen mit Zugang zu Auftragsdaten zur Vertraulichkeit verpflichtet sind und nur die für ihre Aufgabe erforderlichen Berechtigungen erhalten.",
    "SchichtFunk führt die gesetzlich erforderlichen Verzeichnisse und Nachweise, arbeitet mit der zuständigen Aufsichtsbehörde zusammen und stellt dem Kunden die Informationen zur Verfügung, die dieser für den Nachweis der Einhaltung von Artikel 28 DSGVO benötigt.",
    "SchichtFunk verwendet Auftragsdaten nicht für Werbung, den Verkauf von Daten, allgemeine Nutzerprofile oder das Training eigener Modelle. Eine darüber hinausgehende Nutzung setzt eine gesonderte Rechtsgrundlage und transparente Information voraus.",
])

add_numbered(doc, "5", "Technische und organisatorische Maßnahmen", [
    "SchichtFunk setzt die in Anlage 2 und in der beigefügten TOM-Unterlage beschriebenen technischen und organisatorischen Maßnahmen um. Die Maßnahmen berücksichtigen Art, Umfang, Umstände und Zwecke der Verarbeitung sowie die Risiken für die Rechte und Freiheiten der betroffenen Personen.",
    "SchichtFunk darf Maßnahmen technisch weiterentwickeln oder durch gleichwertige Maßnahmen ersetzen, sofern das vereinbarte Schutzniveau nicht unterschritten wird. Wesentliche nachteilige Änderungen teilt SchichtFunk dem Kunden vor ihrem Wirksamwerden mit.",
    "Der Kunde schützt seine Administrationskonten, vergibt Rollen nach dem Erforderlichkeitsprinzip, prüft Berechtigungen regelmäßig und sorgt dafür, dass seine Nutzer sichere Endgeräte und Zugangsdaten verwenden. Die Verantwortungsabgrenzung ist in Anlage 2 festgehalten.",
])

add_numbered(doc, "6", "Unterstützung bei Betroffenenrechten", [
    "SchichtFunk unterstützt den Kunden unter Berücksichtigung der Art der Verarbeitung durch geeignete technische und organisatorische Maßnahmen bei Auskunft, Berichtigung, Löschung, Einschränkung, Datenübertragbarkeit, Widerspruch und der Prüfung automatisierter Entscheidungen.",
    "Geht ein Antrag einer betroffenen Person unmittelbar bei SchichtFunk ein und lässt er sich einem Kunden zuordnen, leitet SchichtFunk ihn unverzüglich an den Kunden weiter. SchichtFunk beantwortet ihn nicht inhaltlich, sofern der Kunde hierzu keine dokumentierte Weisung erteilt hat.",
    "Der Kunde prüft Identität, Zuständigkeit, Rechtsgrundlage und Fristen. Er teilt SchichtFunk die für die technische Unterstützung notwendigen Identifikatoren mit, ohne unnötige Daten zu übermitteln.",
])

add_numbered(doc, "7", "Datenschutz Folgenabschätzung und Aufsicht", [
    "SchichtFunk unterstützt den Kunden mit den verfügbaren Produktinformationen, TOM, Datenflussangaben und Nachweisen bei Datenschutz-Folgenabschätzungen und gegebenenfalls bei einer vorherigen Konsultation der Aufsichtsbehörde.",
    "Für den Einsatz mit Beschäftigtendaten, Krankheitsabwesenheiten, Personalakten, Arbeitszeit- oder QR-Daten prüft der Kunde vor Produktivbeginn, ob eine Datenschutz-Folgenabschätzung erforderlich ist. Die vorhandene SchichtFunk-Schwellenprüfung ersetzt die kundenspezifische Bewertung und Freigabe nicht.",
])

add_numbered(doc, "8", "Sicherheitsvorfälle", [
    "SchichtFunk informiert den Kunden unverzüglich nach Bekanntwerden einer Verletzung des Schutzes von Auftragsdaten. Die Meldung erfolgt an den in Anlage 1 genannten Sicherheitskontakt und enthält die verfügbaren Angaben zu Art und Umfang, betroffenen Daten und Personen, möglichen Folgen sowie getroffenen oder vorgeschlagenen Maßnahmen. Fehlende Angaben werden nachgereicht.",
    "SchichtFunk unterstützt den Kunden bei der Untersuchung, Eindämmung, Dokumentation und Erfüllung der Pflichten nach Artikeln 33 und 34 DSGVO. Die Entscheidung über eine Meldung an die Aufsichtsbehörde oder betroffene Personen trifft der Kunde als Verantwortlicher.",
    "Vereinbarte Reaktionszielzeit für die erste Mitteilung: [[unverzüglich; optionales internes Ziel in Stunden ergänzen]]. Sicherheitskontakt des Kunden: [[E-Mail und Telefon]].",
])

add_numbered(doc, "9", "Unterauftragsverarbeiter", [
    "Der Kunde erteilt SchichtFunk die allgemeine schriftliche Genehmigung, die in Anlage 3 aufgeführten Unterauftragsverarbeiter für die dort beschriebenen Leistungen einzusetzen. SchichtFunk verpflichtet jeden Unterauftragsverarbeiter vertraglich zu einem im Wesentlichen gleichwertigen Datenschutzniveau und bleibt gegenüber dem Kunden für die Erfüllung der Pflichten verantwortlich.",
    "SchichtFunk informiert den Kunden mindestens [[30]] Kalendertage vor der beabsichtigten Hinzuziehung oder Ersetzung eines Unterauftragsverarbeiters über die Änderung. Der Kunde kann innerhalb von [[14]] Kalendertagen aus nachweisbaren datenschutzrechtlichen Gründen widersprechen. Die Parteien suchen zunächst eine zumutbare Lösung. Ist keine Lösung möglich, gelten die Regelungen des Hauptvertrags zur Beendigung der betroffenen Leistung.",
    "Für Unterauftragsverarbeiter, deren eigene aktuelle Unterauftragnehmerlisten Bestandteil ihrer DPA sind, darf SchichtFunk auf die dokumentierte Anbieterquelle verweisen. SchichtFunk führt einen aktuellen Nachweisstand und prüft Änderungen mindestens jährlich sowie anlassbezogen.",
])

add_numbered(doc, "10", "Internationale Datenübermittlungen", [
    "SchichtFunk übermittelt Auftragsdaten nur auf dokumentierte Weisung des Kunden oder aufgrund einer gesetzlichen Verpflichtung in ein Drittland oder an eine internationale Organisation. Jede Übermittlung muss die Anforderungen des Kapitels V DSGVO erfüllen.",
    "Das primäre Supabase-Projekt ist in der Region eu-central-1 in Frankfurt eingerichtet. Support-, Sicherheits- und Unterauftragnehmerzugriffe können dennoch einen Drittlandbezug haben. Soweit erforderlich, werden Angemessenheitsbeschlüsse, EU-Standardvertragsklauseln und ergänzende Maßnahmen nach dem jeweiligen DPA eingesetzt und dokumentiert.",
])

add_numbered(doc, "11", "Löschung Rückgabe und Vertragsende", [
    "Nach Ende der Verarbeitungsleistungen löscht SchichtFunk die Auftragsdaten nach Wahl des Kunden oder stellt sie in einem vereinbarten gängigen Format zur Rückgabe bereit, sofern keine gesetzliche Pflicht zur weiteren Speicherung besteht. Die Auswahl des Kunden und die Exportfrist werden im Offboardingauftrag dokumentiert.",
    "Der Kunde erteilt vor Vertragsende eine dokumentierte Weisung: [[Datenrückgabe und anschließende Löschung / unmittelbare Löschung]]. Vorgesehene Bereitstellungsfrist für einen Export: [[bis zu 30 Tage oder abweichende Frist]].",
    "Die Löschung umfasst Datenbankdaten, Mandantenzuordnungen, Push-Abonnements und dem Kunden zugeordnete Storage-Objekte. Authentifizierungskonten werden nur gelöscht, wenn keine zulässige Zuordnung zu einem anderen Mandanten besteht. Gesetzliche Aufbewahrung, Legal Holds und referenzielle Nachweispflichten können eine Sperrung oder Pseudonymisierung anstelle sofortiger Löschung erfordern.",
    "Backups laufen nach der für den aktiven Anbieterplan geltenden Frist aus und werden nicht zur gewöhnlichen Verarbeitung wiederhergestellt. Wird ein Backup aus Sicherheitsgründen zurückgespielt, sind zwischenzeitlich fällige Löschungen erneut anzuwenden. Die konkrete Fristenmatrix wird für den Kunden in Anlage 1 festgelegt.",
])

add_numbered(doc, "12", "Nachweise und Kontrollen", [
    "SchichtFunk stellt auf Anfrage geeignete Nachweise zur Verfügung, insbesondere die aktuelle TOM-Fassung, das Auftragsverarbeiterregister, vorhandene Prüf- oder Zertifizierungsnachweise der Infrastruktur-Anbieter und dokumentierte Ergebnisse eigener Sicherheits- und Wiederherstellungsprüfungen.",
    "Der Kunde darf die Einhaltung dieser Vereinbarung in angemessenem Umfang selbst oder durch einen zur Vertraulichkeit verpflichteten Prüfer kontrollieren. Prüfungen sollen grundsätzlich durch Dokumentenauskunft erfolgen. Eine weitergehende Prüfung ist zulässig, wenn Dokumente nicht ausreichen, ein konkreter Anlass besteht oder eine Aufsichtsbehörde sie verlangt.",
    "Prüfungen werden rechtzeitig angekündigt, während üblicher Geschäftszeiten durchgeführt und so begrenzt, dass Daten anderer Kunden, Geschäftsgeheimnisse und die Systemsicherheit geschützt bleiben. Bei einem nachgewiesenen wesentlichen Verstoß trägt SchichtFunk die angemessenen Abhilfekosten; im Übrigen richtet sich die Kostenverteilung nach dem Hauptvertrag und dem tatsächlichen Aufwand.",
])

add_numbered(doc, "13", "Vertraulichkeit und Offenlegung", [
    "Beide Parteien behandeln personenbezogene Daten, Sicherheitsinformationen und nicht öffentliche Nachweise vertraulich. Gesetzliche Offenlegungspflichten bleiben unberührt. Soweit zulässig, informiert die offenlegende Partei die andere Partei vorab.",
    "Die Vertraulichkeits- und Datenschutzpflichten gelten nach Ende des Hauptvertrags fort, solange Auftragsdaten oder vertrauliche Informationen vorhanden sind.",
])

add_numbered(doc, "14", "Haftung und Rangfolge", [
    "Für die Haftung gelten Artikel 82 DSGVO, die sonstigen zwingenden gesetzlichen Bestimmungen und die Haftungsregelungen des Hauptvertrags. Diese Vereinbarung schränkt Rechte betroffener Personen oder Befugnisse von Aufsichtsbehörden nicht ein.",
    "Bei einem Widerspruch zwischen dieser Vereinbarung und dem Hauptvertrag hat diese Vereinbarung für Fragen der Auftragsverarbeitung Vorrang. Zwingende gesetzliche Regelungen und wirksam einbezogene EU-Standardvertragsklauseln gehen vor.",
])

add_numbered(doc, "15", "Laufzeit Änderungen und Schlussbestimmungen", [
    "Diese Vereinbarung tritt am [[Datum]] in Kraft. Sie endet erst, wenn SchichtFunk alle Auftragsdaten vertragsgemäß gelöscht oder zurückgegeben hat.",
    "Änderungen und Ergänzungen bedürfen der Textform, soweit keine strengere Form vorgeschrieben ist. Dies gilt auch für die Änderung dieser Formregelung. Elektronische Signaturen und elektronisch dokumentierte Annahmen sind zulässig.",
    "Sollte eine Bestimmung unwirksam sein oder werden, bleiben die übrigen Bestimmungen unberührt. Die Parteien ersetzen die unwirksame Bestimmung durch eine wirksame Regelung, die dem datenschutzrechtlichen Zweck möglichst nahekommt.",
    "Es gelten das im Hauptvertrag vereinbarte Recht und der dort vereinbarte Gerichtsstand, soweit zwingendes Datenschutzrecht nichts anderes bestimmt.",
])

doc.add_heading("Unterzeichnung", level=1)
add_table(
    doc,
    ["Für den Kunden", "Für SchichtFunk"],
    [
        ("Ort und Datum\n\n[[Ort und Datum]]", "Ort und Datum\n\n[[Ort und Datum]]"),
        ("Name und Funktion\n\n[[Name und Funktion]]", "Name und Funktion\n\nAlexander Kirchner"),
        ("Unterschrift oder elektronische Annahme\n\n\n", "Unterschrift oder elektronische Annahme\n\n\n"),
    ],
    widths=[3.35, 3.35],
)

doc.add_page_break()
doc.add_heading("Anlage 1 Beschreibung der Verarbeitung", level=1)
add_paragraph(doc, "Diese Anlage ist für jeden Kundenmandanten zu vervollständigen und gemeinsam mit der Vereinbarung zu speichern.")

add_table(
    doc,
    ["Angabe", "Kundenbezogene Festlegung"],
    [
        ("Kundenmandant", "[[Name und interne Mandantenkennung]]"),
        ("Vertragsbeginn", "[[Datum]]"),
        ("Vertragsende", "nach Hauptvertrag; konkretes Ende [[Datum oder offen]]"),
        ("Weisungsberechtigte", "[[Name Funktion E-Mail]]"),
        ("Datenschutzkontakt", "[[Name Funktion E-Mail Telefon]]"),
        ("Sicherheitskontakt", "[[Name Funktion E-Mail Telefon]]"),
        ("Betriebsrat oder Personalvertretung", "[[vorhanden / nicht vorhanden / Prüfung offen]]"),
        ("Anzahl Beschäftigte und Standorte", "[[Anzahl und Standorte]]"),
        ("Produktivbeginn", "[[Datum nach Freigabe]]"),
    ],
    widths=[2.2, 4.5],
)

doc.add_heading("Zwecke und Funktionen", level=2)
add_paragraph(doc, "Zweck ist die Bereitstellung und der sichere Betrieb der vom Kunden ausgewählten SchichtFunk-Funktionen. Nicht benötigte Funktionen sind im Kundenmandanten zu deaktivieren oder organisatorisch auszuschließen.")
functions = [
    "Benutzerkonten, Einladungen, Rollen und Berechtigungen",
    "Mitarbeiterstammdaten und Qualifikationen",
    "Dienstplanung, Verfügbarkeiten, Schichtänderungen, Tausch und Marktplatz",
    "Arbeitszeit, Pausen, QR-Buchungen, Stundenkonto und Monatsabschluss",
    "Abwesenheiten einschließlich des notwendigen Status einer Krankheitsabwesenheit",
    "Digitale Personalakte und Dokumente, sofern vom Kunden freigegeben",
    "Benachrichtigungen und freiwillige Web-Push-Zustellung",
    "DATEV-LODAS-Vorbereitung und fachliche Exportnachweise",
    "Audit, Sicherheit, Support und Störungsbearbeitung",
]
for item in functions:
    add_bullet(doc, f"☐ {item}")

doc.add_heading("Art und Häufigkeit", level=2)
add_paragraph(doc, "Erheben, Erfassen, Ordnen, Speichern, Abfragen, Anzeigen, Ändern, Übermitteln innerhalb der genehmigten Verarbeitungskette, Einschränken, Exportieren, Pseudonymisieren und Löschen. Die Verarbeitung erfolgt fortlaufend während der Nutzung und anlassbezogen bei Support, Export, Sicherheits- oder Löschvorgängen.")

doc.add_heading("Betroffene Personen", level=2)
for item in [
    "Beschäftigte, ehemalige Beschäftigte und Bewerber, soweit der Kunde entsprechende Daten zulässig einstellt",
    "Inhaber, Führungskräfte, Personalverantwortliche, Planer, Administratoren und sonstige berechtigte Nutzer",
    "Kundenkontaktpersonen, Weisungsberechtigte und Datenschutzkontakte",
    "gegebenenfalls Notfallkontakte, sofern der Kunde diese Funktion nutzt",
]:
    add_bullet(doc, item)

doc.add_heading("Datenkategorien", level=2)
add_table(
    doc,
    ["Kategorie", "Typische Inhalte und Begrenzung"],
    [
        ("Konten und Kontakte", "Name, geschäftliche oder private Kontaktangaben, Authentifizierungskennung, Rollen, Status, Einladungsdaten"),
        ("Beschäftigtenstamm", "Personalnummer, Anschrift, Geburts- und Beschäftigungsdaten, Arbeitszeitmodell, Qualifikationen, organisatorische Zuordnung"),
        ("Planung", "Schichten, Einsatzorte, Verfügbarkeit, Zuordnung, Bestätigung, Tausch- und Änderungsanträge"),
        ("Zeit und Entgeltvorbereitung", "Arbeitsbeginn und -ende, Pausen, QR-Buchung, Stundenkonto, Zuschlags- und DATEV-Vorbereitungsdaten"),
        ("Abwesenheiten", "Zeitraum, Kategorie und Freigabestatus; Krankheitsstatus als Gesundheitsdatum; Diagnosen sind nicht regulär vorgesehen"),
        ("Personalakte", "vom Kunden freigegebene Dokumente, Qualifikations- und Fristangaben; Inhalt und Rechtsgrundlage legt der Kunde fest"),
        ("Benachrichtigungen", "fachliche Nachricht, Zielansicht, Zustellstatus, Push-Endpunkt und Gerätemetadaten; Nachrichten sind datensparsam zu formulieren"),
        ("Audit und Sicherheit", "Nutzer- und Mandantenkennung, Zeitstempel, Aktion, technische Metadaten, Sicherheits- und Fehlerdaten"),
        ("Support", "Kontakt, Anfrageinhalt und notwendige technische Diagnose; Echtdaten nur soweit erforderlich und geschützt"),
    ],
    widths=[1.75, 4.95],
)

doc.add_heading("Besondere Kategorien und ausgeschlossene Verarbeitungen", level=2)
add_paragraph(doc, "Vereinbart ist ausschließlich die für die Abwesenheitsverwaltung erforderliche Information, dass eine Person krankheitsbedingt abwesend ist, sowie weitere besondere Kategorien nur nach ausdrücklicher Ergänzung: [[Angabe oder keine]]. Diagnosen, Behandlungsunterlagen, biometrische Identifikation, GPS-Bewegungsprofile und verdecktes Leistungsscoring sind nicht Bestandteil der Standardleistung.")

doc.add_heading("Löschfristen und Rückgabe", level=2)
add_paragraph(doc, "Der Kunde ergänzt vor Produktivbeginn seine fachlich und rechtlich geprüften Fristen. Bis zur Freigabe gelten keine pauschalen automatischen Löschfristen für Beschäftigtendaten.")
add_table(
    doc,
    ["Datenbereich", "Kundenfrist oder Ereignis", "Aktion"],
    [
        ("Einladungen und technische Kurzzeitdaten", "[[Frist]]", "löschen"),
        ("Push-Abonnements und Benachrichtigungen", "[[Frist / bei Abmeldung]]", "löschen"),
        ("Dienstplan und Arbeitszeitnachweise", "[[Frist nach Rechtsprüfung]]", "sperren, exportieren oder löschen"),
        ("Abwesenheiten und Krankheitsstatus", "[[Frist nach Rechtsprüfung]]", "löschen oder erforderlichenfalls sperren"),
        ("Personalakten-Dokumente", "[[dokumenttypbezogene Frist]]", "Storage-Objekt und Metadaten koordiniert löschen"),
        ("DATEV- und Monatsnachweise", "[[Frist nach Steuer- und Handelsrecht]]", "aufbewahren, anschließend löschen oder anonymisieren"),
        ("Audit- und Incidentnachweise", "[[Frist]]", "inhaltlich minimieren, anschließend löschen oder anonymisieren"),
        ("Vertragsende", "[[Exportfrist, höchstens vereinbarter Zeitraum]]", "Rückgabe und bestätigte Löschung"),
    ],
    widths=[1.7, 2.35, 2.65],
)

doc.add_page_break()
doc.add_heading("Anlage 2 Technische und organisatorische Maßnahmen", level=1)
add_paragraph(doc, "Vertragsbestandteil ist die separat bereitgestellte Unterlage SchichtFunk Technische und organisatorische Maßnahmen in der bei Vertragsschluss bestätigten Fassung [[Versionsnummer und Datum]]. Die folgenden Maßnahmen beschreiben den derzeit vorgesehenen Mindeststandard.")
add_table(
    doc,
    ["Bereich", "Mindestmaßnahmen"],
    [
        ("Organisation und Vertraulichkeit", "Rollen- und Weisungskonzept; Vertraulichkeitsbindung; keine Geheimnisse oder Echtdaten im Quellcode; dokumentierte Anbieter-DPA und Unterauftragnehmer"),
        ("Authentifizierung", "Supabase Auth; gehashte Passwörter; Leaked Password Protection; MFA-Unterstützung; zeitlich begrenzte Sitzungen; keine Service-Schlüssel im Browser"),
        ("Berechtigungen und Mandantentrennung", "Rollen nach Erforderlichkeitsprinzip; Row Level Security; company_id und aktive Mitgliedschaft; serverseitige Prüfung privilegierter Funktionen"),
        ("Personalakte und Storage", "privater Bucket; rollenbeschränkte Funktionen; mandantenbezogene Objektpfade; kurzlebige signierte Download-URLs"),
        ("Transport und Webschutz", "TLS; HSTS; Content Security Policy; Schutz vor Einbettung und MIME-Sniffing; keine API- oder QR-Antworten im Service-Worker-Cache"),
        ("Protokollierung", "mandantenbezogene Auditereignisse für sicherheits- und fachlich relevante Vorgänge; begrenzte technische Anbieterlogs; kein öffentlicher Marketingtracker im geprüften Stand"),
        ("Verfügbarkeit", "Supabase Pro in eu-central-1 Frankfurt; tägliche Datenbankbackups nach aktivem Tarif; getesteter logischer DB- und privater Storage-Wiederherstellungsweg; Wiederanlaufplanung"),
        ("Löschung", "kundenspezifische Fristen; Legal Holds; Vier-Augen-Freigabe für Löschaufträge; koordinierte Sperr-, Export-, Storage-, Datenbank-, Push- und Auth-Schritte"),
        ("Entwicklung und Änderung", "Versionsverwaltung; automatisierte Tests; getrennte Vorschau vor Domainumschaltung; Geheimnisse außerhalb des Repository; dokumentierte Freigabe und Rückfallplanung"),
        ("Prüfung", "regelmäßige RLS- und RPC-Prüfungen; Mandantentrennungstests; jährliche sowie anlassbezogene Überprüfung von TOM, Restore, Unterauftragnehmern und Risiken"),
    ],
    widths=[1.7, 5.0],
)

doc.add_heading("Verantwortung des Kunden", level=2)
for item in [
    "Konten ausschließlich personenbezogen vergeben und gemeinsame Administrationskonten vermeiden",
    "Rollen, Austritte und Berechtigungen regelmäßig prüfen",
    "MFA für privilegierte Nutzer nach Freigabe aktivieren und sichere Endgeräte verlangen",
    "Diagnosen und unnötige Freitexte organisatorisch ausschließen",
    "fachliche Löschfristen, Rechtsgrundlagen und Legal Holds festlegen",
    "Beschäftigte informieren und Betriebsrat oder Personalvertretung einbeziehen",
    "Exporte und heruntergeladene Dateien außerhalb von SchichtFunk angemessen schützen und löschen",
]:
    add_bullet(doc, item)

doc.add_heading("Vor der Unterzeichnung zu bestätigen", level=2)
for item in [
    "☐ aktuelle TOM-Fassung beigefügt und geprüft",
    "☐ kundenspezifische DSFA und Beteiligungspflichten geprüft",
    "☐ Weisungs-, Datenschutz- und Sicherheitskontakte benannt",
    "☐ Lösch- und Aufbewahrungsfristen freigegeben",
    "☐ erlaubte Funktionen und besondere Datenkategorien festgelegt",
    "☐ offene Härtungs- oder Betriebsnachweise bewertet und akzeptiert",
]:
    add_bullet(doc, item)

doc.add_page_break()
doc.add_heading("Anlage 3 Genehmigte Unterauftragsverarbeiter", level=1)
add_paragraph(doc, "Stand der Vorlage: 13.09.2026. Vor Vertragsabschluss ist die aktuelle Anbieter- und Unterauftragnehmerlage erneut zu prüfen.")
add_table(
    doc,
    ["Unterauftragsverarbeiter", "Leistung und Datenbezug", "Ort und Garantien"],
    [
        (
            "IONOS SE",
            "Primäres statisches Hosting der Webseite und PWA über Deploy Now; technische Verbindungs-, Sicherheits- und Logdaten. Keine Beschäftigten- oder Arbeitszeitdaten im statischen Build.",
            "Deutschland und weitere in den IONOS-Vertragsanlagen benannte Orte; IONOS-AVV und TOM; etwaige weitere Unterauftragnehmer gemäß aktueller IONOS-Anlage.",
        ),
        (
            "Supabase Pte. Ltd",
            "PostgreSQL-Datenbank, Authentifizierung, Realtime, Storage, Edge Functions, Logs und Backups für die SchichtFunk-Anwendungsdaten.",
            "Vertragspartner in Singapur; primäre Projektdaten in eu-central-1 Frankfurt; DPA, EU-Standardvertragsklauseln, TIA und veröffentlichte Unterauftragnehmerliste.",
        ),
    ],
    widths=[1.45, 2.85, 2.4],
)

doc.add_heading("Anbieterquellen", level=2)
for source in [
    "IONOS AVV und Anlagen: https://www.ionos.de/terms-gtc/avv/",
    "Supabase DPA: https://supabase.com/legal/customer-resources/data-processing-addendum",
    "Supabase Unterauftragnehmerliste: https://supabase.com/legal/customer-resources/subprocessor-list",
    "SchichtFunk Nachweisregister: documentation/avv-dpa-evidence-register-2026-09-12.md",
]:
    add_bullet(doc, source)

doc.add_heading("Nicht freigegebene Rückfallumgebung", level=2)
add_paragraph(doc, "Vercel bleibt technisch als Rückfallstand vorhanden, ist im aktuellen Hobby-Tarif jedoch nicht für einen kommerziellen Produktivbetrieb freigegeben. Vercel darf für Auftragsdaten dieses Kunden erst nach geeigneter kommerzieller Vertrags- und DPA-Grundlage, Aktualisierung dieser Anlage und dokumentierter Kundeninformation aktiviert werden.")

doc.add_heading("Web Push", level=2)
add_paragraph(doc, "Web Push ist für diesen Kunden [[aktiviert / deaktiviert]]. Bei Aktivierung wird die Nachricht über den vom Endgerät gewählten Browser- oder Betriebssystem-Pushdienst zugestellt. Der konkrete Dienst kann von Gerät und Browser abhängen. Push-Inhalte sind auf neutrale, datensparsame Hinweise zu begrenzen; Gesundheits- oder Personalaktendetails dürfen nicht im Nachrichtentext erscheinen.")

doc.add_heading("Änderungsmitteilungen", level=2)
add_paragraph(doc, "Mitteilungen über Änderungen der Unterauftragsverarbeiter werden an [[Kunden-E-Mail]] gesendet. Ein Widerspruch ist innerhalb der in Ziffer 9 vereinbarten Frist an info@schichtfunk.de zu richten.")

doc.add_page_break()
doc.add_heading("Anlage 4 Kontakte und Weisungsberechtigungen", level=1)
add_table(
    doc,
    ["Funktion", "Name", "Kontakt", "Berechtigungsumfang"],
    [
        ("Vertragskontakt", "[[Name]]", "[[E-Mail Telefon]]", "Hauptvertrag und AVV"),
        ("Weisung", "[[Name]]", "[[E-Mail Telefon]]", "fachliche Weisungen und Exporte"),
        ("Datenschutz", "[[Name]]", "[[E-Mail Telefon]]", "Betroffenenrechte, DSFA, Behörden"),
        ("Sicherheit", "[[Name]]", "[[E-Mail Telefon]]", "Sicherheitsvorfälle und Sofortmaßnahmen"),
        ("Löschfreigabe 1", "[[Name]]", "[[E-Mail Telefon]]", "Antrag oder Freigabe"),
        ("Löschfreigabe 2", "[[Name]]", "[[E-Mail Telefon]]", "getrennte Gegenfreigabe"),
    ],
    widths=[1.55, 1.2, 1.65, 2.3],
)

doc.add_heading("Zulässige Weisungskanäle", level=2)
for item in [
    "☐ Administrationsfunktionen im SchichtFunk-Kundenmandanten",
    "☐ E-Mail von einer oben benannten Adresse an info@schichtfunk.de",
    "☐ signiertes Dokument oder vereinbartes Ticketsystem [[Angabe]]",
    "☐ Notfallkontakt [[Telefonnummer und Identitätsprüfung]]",
]:
    add_bullet(doc, item)

doc.add_heading("Quellen und Einordnung", level=1)
add_paragraph(doc, "Diese Vorlage orientiert sich an Artikel 28 DSGVO und an den von der Europäischen Kommission mit Durchführungsbeschluss EU 2021/915 veröffentlichten Standardvertragsklauseln zwischen Verantwortlichen und Auftragsverarbeitern. Sie übernimmt die gesetzlichen Pflichtinhalte in einer auf SchichtFunk zugeschnittenen Struktur und ist nicht als unveränderte amtliche Standardklausel ausgewiesen.")
for source in [
    "Datenschutz-Grundverordnung, insbesondere Artikel 28 bis 36: https://eur-lex.europa.eu/eli/reg/2016/679/deu",
    "Durchführungsbeschluss EU 2021/915: https://eur-lex.europa.eu/eli/dec_impl/2021/915/oj?locale=de",
    "BfDI Information zur Auftragsverarbeitung: https://www.bfdi.bund.de/SharedDocs/Publikationen/Infobroschueren/INFO6.html",
]:
    add_bullet(doc, source)

doc.core_properties.title = "Vereinbarung zur Auftragsverarbeitung für SchichtFunk Kunden"
doc.core_properties.subject = "Vertragsvorlage nach Artikel 28 DSGVO"
doc.core_properties.author = "Alexander Kirchner / SchichtFunk"
doc.core_properties.keywords = "SchichtFunk, AVV, Auftragsverarbeitung, DSGVO, Artikel 28"

doc.save(OUTPUT)
print(OUTPUT)
